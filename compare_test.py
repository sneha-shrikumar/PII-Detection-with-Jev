"""Final test: Jev + Haiku vs Haiku alone vs Microsoft Presidio, scored with the same release gate.

Each system is run once and saved to runs/final_<system>.json, then `compare` scores all three:
    python compare_test.py jev        Jev + Haiku pipeline (pii_detect.py, v2 questions, traced to LangSmith)
    python compare_test.py haiku      Claude Haiku 4.5 alone on every chat: same prompt minus the screening hint,
                                      same verbatim check
    python compare_test.py presidio   Presidio analyzer (spaCy en_core_web_lg), default settings, runs locally
    python compare_test.py colab      or, where spaCy can't run locally: writes presidio_colab.ipynb for Google Colab,
    python compare_test.py presidio-import results/final_presidio_colab.json   then reads back what it downloads
    python compare_test.py compare    scores the three side by side -> runs/final_comparison_<time>.xlsx
    python compare_test.py redact     redacted chats, Jev + Haiku vs Haiku alone -> runs/redacted_test_<time>.json/.xlsx

The release gate is applied per expected PII span, the same way for all three systems:
    SAFE  the span was extracted (so it can be redacted), or the chat went to a person for that category
    LEAK  anything else
plus zero invented text, cost per 1,000 chats, and precision: over 90% of extracted items must be real PII.
Only the Jev pipeline can send chats to a person; Haiku alone and Presidio have no "unsure" signal.
"""
import datetime as dt
import json
import sys
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import pii_detect as P

TEST = "ecom_realistic_test_350.csv"
RUNS = P.HERE / "runs"
SYSTEMS = {"jev": "Jev + Haiku", "haiku": "Haiku alone", "presidio": "Presidio"}

# Presidio entity -> our category. Presidio has no address or date-of-birth entity, so the
# nearest ones (LOCATION, DATE_TIME) stand in; ID entities from any country count as NRIC_PASSPORT.
PRESIDIO_MAP = {
    "PERSON": "FULL_NAME", "CREDIT_CARD": "CARD_NUMBER", "PHONE_NUMBER": "PHONE", "EMAIL_ADDRESS": "EMAIL",
    "LOCATION": "ADDRESS", "DATE_TIME": "DOB",
    "SG_NRIC_FIN": "NRIC_PASSPORT", "US_PASSPORT": "NRIC_PASSPORT", "IN_PASSPORT": "NRIC_PASSPORT",
    "US_SSN": "NRIC_PASSPORT", "IN_AADHAAR": "NRIC_PASSPORT", "IN_PAN": "NRIC_PASSPORT",
    "UK_NHS": "NRIC_PASSPORT", "UK_NINO": "NRIC_PASSPORT", "AU_TFN": "NRIC_PASSPORT",
    "ES_NIE": "NRIC_PASSPORT", "ES_NIF": "NRIC_PASSPORT", "IT_IDENTITY_CARD": "NRIC_PASSPORT",
    "IT_PASSPORT": "NRIC_PASSPORT", "PL_PESEL": "NRIC_PASSPORT", "FI_PERSONAL_IDENTITY_CODE": "NRIC_PASSPORT",
    "KR_RRN": "NRIC_PASSPORT", "TH_TNIN": "NRIC_PASSPORT",
}


def load_rows():
    import csv
    with open(P.HERE / TEST, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def save(name, results):
    RUNS.mkdir(exist_ok=True)
    for r in results:
        if r.get("labels"):
            r["labels"] = dict(r["labels"], types=sorted(r["labels"]["types"]))
    path = RUNS / f"final_{name}.json"
    path.write_text(json.dumps(results, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"saved {path}")


def load(name):
    results = json.loads((RUNS / f"final_{name}.json").read_text(encoding="utf-8"))
    for r in results:
        if r.get("labels"):
            r["labels"]["types"] = set(r["labels"]["types"])
    return results


# ---------------------------------------------------------------- 1. Jev + Haiku (the normal pipeline)
def run_jev():
    kept = {}
    real_run = P.run

    def run_and_keep(*a, **k):
        kept["results"] = real_run(*a, **k)
        return kept["results"]

    P.run = run_and_keep
    sys.argv = ["pii_detect.py", TEST]
    P.main()
    for r in kept["results"]:
        r.pop("raw_answers", None)
    save("jev", kept["results"])


# ---------------------------------------------------------------- 2. Haiku alone
def run_haiku():
    key = P.env_key("OPENROUTER_API_KEY")
    rows = load_rows()
    # No screening model, so the same prompt without the hint rule, and an empty hint. (Telling Haiku that
    # all 7 categories were flagged made it return lone first names as FULL_NAME in 48 clean chats.)
    hint_rule = P.LLM_SYSTEM.rsplit("\n- ", 1)[1]
    assert hint_rule.startswith("The hint lists"), hint_rule
    P.LLM_SYSTEM = P.LLM_SYSTEM.replace(hint_rule, "Return an empty list if there is no PII.")
    every = {"types_found": [], "unsure": []}

    def one(row):
        try:
            x = P.extract_pii(row["text"], every, key)
            return {"id": row["id"], "text": row["text"], "labels": P.read_labels(row), "error": "",
                    "extracted": x["items"], "rejected": x["rejected"], "cost_usd": x["cost"],
                    "tokens_in": x["input_tokens"], "tokens_out": x["output_tokens"], "ms": x["latency_ms"]}
        except Exception as e:
            print(f"  FAILED {row['id']}: {e}")
            return {"id": row["id"], "text": row["text"], "labels": P.read_labels(row), "error": str(e)[:500]}

    print(f"Haiku alone: {len(rows)} chats")
    results = []
    with ThreadPoolExecutor(max_workers=10) as pool:
        for i, fut in enumerate(as_completed([pool.submit(one, r) for r in rows]), 1):
            results.append(fut.result())
            if i % 50 == 0 or i == len(rows):
                print(f"  {i}/{len(rows)} done")
    results.sort(key=lambda r: r["id"])
    save("haiku", results)


# ---------------------------------------------------------------- 3. Presidio
def run_presidio():
    from presidio_analyzer import AnalyzerEngine
    from presidio_analyzer.recognizer_registry import RecognizerRegistry

    registry = RecognizerRegistry()
    registry.load_predefined_recognizers()
    engine = AnalyzerEngine(registry=registry)
    supported = set(engine.get_supported_entities("en"))
    entities = sorted(e for e in PRESIDIO_MAP if e in supported)
    print(f"Presidio: entities used {entities}")
    print(f"          not available in this install: {sorted(set(PRESIDIO_MAP) - supported)}")

    results = []
    for row in load_rows():
        text = row["text"]
        start = time.perf_counter()
        found = engine.analyze(text=text, language="en", entities=entities)
        ms = (time.perf_counter() - start) * 1000
        items, seen = [], set()
        for f in sorted(found, key=lambda f: f.start):
            it = (PRESIDIO_MAP[f.entity_type], text[f.start:f.end])
            if it not in seen:
                seen.add(it)
                items.append({"label": it[0], "text": it[1], "presidio": f.entity_type, "score": round(f.score, 2)})
        results.append({"id": row["id"], "text": text, "labels": P.read_labels(row), "error": "",
                        "extracted": items, "rejected": [], "cost_usd": 0.0, "ms": ms})
    save("presidio", results)


def import_presidio(path):
    """Presidio run in Google Colab (presidio_colab.ipynb): add the text and labels back, save as final_presidio."""
    found = {r["id"]: r for r in json.loads(Path(path).read_text(encoding="utf-8"))}
    rows = load_rows()
    missing = [r["id"] for r in rows if r["id"] not in found]
    if missing:
        sys.exit(f"{len(missing)} chats missing from {path}, e.g. {missing[:3]}")
    results = []
    for row in rows:
        x = found[row["id"]]
        items = [i for i in x["extracted"] if i["text"] in row["text"]]
        assert len(items) == len(x["extracted"]), f"{row['id']}: Presidio span not in the chat"
        results.append({"id": row["id"], "text": row["text"], "labels": P.read_labels(row), "error": "",
                        "extracted": items, "rejected": [], "cost_usd": 0.0, "ms": x["ms"]})
    print(f"Presidio (Colab) {x.get('versions', '')}: {len(results)} chats")
    save("presidio", results)


def build_colab_notebook():
    """Writes presidio_colab.ipynb with the test chats (id and text only, no labels) built in."""
    chats = [{"id": r["id"], "text": r["text"]} for r in load_rows()]
    code = lambda src: {"cell_type": "code", "metadata": {}, "execution_count": None, "outputs": [], "source": src}
    nb = {"nbformat": 4, "nbformat_minor": 5,
          "metadata": {"kernelspec": {"name": "python3", "display_name": "Python 3"}, "colab": {"provenance": []}},
          "cells": [
              {"cell_type": "markdown", "metadata": {}, "source":
               "# Presidio baseline for the PII final test\n\nRuntime > Run all. The last cell downloads "
               "`final_presidio_colab.json`; move it into the project folder and run\n"
               "`python compare_test.py presidio-import final_presidio_colab.json` then `python compare_test.py compare`.\n\n"
               "If Colab asks to restart the session after the install, restart it and Run all again."},
              code("!pip install -q presidio-analyzer\n!python -m spacy download en_core_web_lg -q"),
              code("CHATS = " + json.dumps(chats, ensure_ascii=False, indent=0)),
              code("PRESIDIO_MAP = " + json.dumps(PRESIDIO_MAP, indent=1)),
              code('''import json, time
from importlib.metadata import version
from presidio_analyzer import AnalyzerEngine
from presidio_analyzer.recognizer_registry import RecognizerRegistry

registry = RecognizerRegistry()
registry.load_predefined_recognizers()
engine = AnalyzerEngine(registry=registry)  # default: spaCy en_core_web_lg, default score threshold
supported = set(engine.get_supported_entities("en"))
entities = sorted(e for e in PRESIDIO_MAP if e in supported)
versions = f"presidio-analyzer {version('presidio-analyzer')}, spacy {version('spacy')}"
print(versions)
print("entities used:", entities)
print("not available:", sorted(set(PRESIDIO_MAP) - supported))

out = []
for c in CHATS:
    text = c["text"]
    t0 = time.perf_counter()
    found = engine.analyze(text=text, language="en", entities=entities)
    ms = (time.perf_counter() - t0) * 1000
    items, seen = [], set()
    for f in sorted(found, key=lambda f: f.start):
        it = (PRESIDIO_MAP[f.entity_type], text[f.start:f.end])
        if it not in seen:
            seen.add(it)
            items.append({"label": it[0], "text": it[1], "presidio": f.entity_type, "score": round(f.score, 2)})
    out.append({"id": c["id"], "extracted": items, "ms": ms, "versions": versions})
print(len(out), "chats,", sum(len(o["extracted"]) for o in out), "items found")'''),
              code('''with open("final_presidio_colab.json", "w", encoding="utf-8") as fh:
    json.dump(out, fh, ensure_ascii=False)
from google.colab import files
files.download("final_presidio_colab.json")'''),
          ]}
    path = P.HERE / "presidio_colab.ipynb"
    path.write_text(json.dumps(nb, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"saved {path}")


# ---------------------------------------------------------------- 4. compare
def gate(results, system):
    """Span-level release gate, identical for every system."""
    ok = [r for r in results if not r.get("error")]
    rows, high_leaks, leaked_ids = [], 0, {}
    for c in P.CATEGORIES:
        spans = safe_ext = safe_person = 0
        ids = []
        for r in ok:
            gold = [g for g in r["labels"]["spans"] if g["label"] == c]
            if not gold:
                continue
            hits = P.span_match(r["text"], r["extracted"], gold)[0]
            sent = c in r.get("unsure", [])
            for h in hits:
                spans += 1
                if h:
                    safe_ext += 1
                elif sent:
                    safe_person += 1
                else:
                    ids.append(r["id"])
        leaks = spans - safe_ext - safe_person
        if P.RISK[c] == "high":
            high_leaks += leaks
        leaked_ids[c] = sorted(set(ids))
        rows.append({"cat": c, "risk": P.RISK[c], "spans": spans, "extracted": safe_ext,
                     "person": safe_person, "leaks": leaks})
    return rows, high_leaks, leaked_ids


def detection(results):
    """Chat-level yes/no per category. Jev+Haiku uses Jev's answer; the others use 'extracted anything'."""
    ok = [r for r in results if not r.get("error")]
    out = {}
    for c in P.CATEGORIES:
        tp = fp = fn = 0
        for r in ok:
            said = c in r["types_found"] if "types_found" in r else any(i["label"] == c for i in r["extracted"])
            truth = c in r["labels"]["types"]
            tp, fp, fn = tp + (said and truth), fp + (said and not truth), fn + (truth and not said)
        out[c] = (tp, fp, fn) + P.prf(tp, fp, fn)
    return out


def spans(results):
    g_hit = g_n = p_hit = p_n = l_ok = l_n = 0
    for r in results:
        if r.get("error"):
            continue
        gh, ph, lo = P.span_match(r["text"], r["extracted"], r["labels"]["spans"])
        g_hit, g_n, p_hit, p_n, l_ok, l_n = g_hit + sum(gh), g_n + len(gh), p_hit + sum(ph), p_n + len(ph), l_ok + sum(lo), l_n + len(lo)
    return {"recall": g_hit / g_n, "precision": p_hit / p_n if p_n else float("nan"),
            "label_acc": l_ok / l_n if l_n else float("nan"), "found": f"{g_hit}/{g_n}", "correct": f"{p_hit}/{p_n}"}


def compare():
    data = {k: load(k) for k in SYSTEMS if (RUNS / f"final_{k}.json").exists()}
    missing = [SYSTEMS[k] for k in SYSTEMS if k not in data]
    if missing:
        print(f"Not run yet, left out: {', '.join(missing)}")
    stats = {}
    for k, res in data.items():
        ok = [r for r in res if not r.get("error")]
        g_rows, high_leaks, leaked_ids = gate(res, k)
        invented = sum(len(r["rejected"]) for r in ok)
        cost_1k = sum(r["cost_usd"] for r in ok) / len(ok) * 1000
        review = sum(bool(r.get("unsure")) for r in ok)
        hit, n = P.item_precision(ok)
        stats[k] = dict(precision=(hit / n if n else float("nan"), hit, n),n=len(ok), failed=len(res) - len(ok), gate=g_rows, high_leaks=high_leaks, leaked_ids=leaked_ids,
                        invented=invented, cost_1k=cost_1k, review=review, det=detection(res), spans=spans(res),
                        clean_flagged=sum(bool(r["extracted"]) and not r["labels"]["has_pii"] for r in ok),
                        clean=sum(not r["labels"]["has_pii"] for r in ok))
    haiku_cost = stats["haiku"]["cost_1k"]

    def verdicts(s, k):
        v = {"1. High-risk leaks (target 0)": (s["high_leaks"], s["high_leaks"] == 0),
             f"2. Precision, all categories (over {P.PRECISION_MIN:.0%})":
                 ("{:.1%} ({}/{})".format(*s["precision"]), s["precision"][0] > P.PRECISION_MIN),
             "3. Invented text (target 0)": (s["invented"], s["invented"] == 0),
             f"4. Cost per 1,000 chats (under ${P.COST_CAP_PER_1K:.2f})": (f"${s['cost_1k']:.3f}", s["cost_1k"] < P.COST_CAP_PER_1K)}
        if k != "haiku":
            v["Cheaper than Haiku alone"] = (f"${s['cost_1k']:.3f} vs ${haiku_cost:.3f}", s["cost_1k"] < haiku_cost)
        return v

    names = list(data)
    tables = []
    gate_rows = []
    for check in verdicts(stats["jev"], "jev"):
        row = [check]
        for k in names:
            v = verdicts(stats[k], k).get(check)
            row.append("-" if v is None else f"{v[0]}  {'PASS' if v[1] else 'FAIL'}")
        gate_rows.append(row)
    overall = ["OVERALL"] + ["PASS" if all(ok for _, ok in verdicts(stats[k], k).values()) else "FAIL" for k in names]
    tables.append(("Release gate (same rules for all three)", ["Check"] + [SYSTEMS[k] for k in names], gate_rows + [overall]))

    rows = []
    for i, c in enumerate(P.CATEGORIES):
        row = [c, P.RISK[c], stats["jev"]["gate"][i]["spans"]]
        for k in names:
            g = stats[k]["gate"][i]
            row += [g["extracted"], g["person"], g["leaks"]]
        rows.append(row)
    cols = ["Category", "Risk", "Expected spans"]
    for k in names:
        cols += [f"{SYSTEMS[k]}: extracted", f"{SYSTEMS[k]}: sent to person", f"{SYSTEMS[k]}: LEAKS"]
    tables.append(("Leaks per expected PII span (safe = extracted, or sent to a person)", cols, rows))

    tables.append(("Leaked chats", ["System"] + P.CATEGORIES,
                   [[SYSTEMS[k]] + [", ".join(stats[k]["leaked_ids"][c]) or "-" for c in P.CATEGORIES] for k in names]))

    tables.append(("Extraction vs expected spans", ["System", "Span recall", "Span precision", "Label accuracy",
                                                    "Expected spans found", "Extracted items correct"],
                   [[SYSTEMS[k], s["recall"], s["precision"], s["label_acc"], s["found"], s["correct"]]
                    for k in names for s in [stats[k]["spans"]]]))

    rows = []
    for c in P.CATEGORIES:
        row = [c]
        for k in names:
            tp, fp, fn, p, r, f = stats[k]["det"][c]
            row += [p, r, f]
        rows.append(row)
    cols = ["Category"] + [f"{SYSTEMS[k]}: {m}" for k in names for m in ("precision", "recall", "F1")]
    tables.append(("Per category, per chat (does the chat contain it?)", cols, rows))

    tables.append(("Workload, false alarms and cost", ["System", "Chats", "Failed", "Sent to a person",
                                                       "Clean chats with something extracted", "Cost per 1,000 chats $"],
                   [[SYSTEMS[k], s["n"], s["failed"], s["review"], f"{s['clean_flagged']}/{s['clean']}", s["cost_1k"]]
                    for k in names for s in [stats[k]]]))

    for title, cols, rows in tables:
        print(f"\n{title}\n  " + " | ".join(cols))
        for row in rows:
            print("  " + " | ".join(f"{v:.3f}" if isinstance(v, float) else str(v) for v in row))
    write_xlsx(tables, data)


def write_xlsx(tables, data):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    bold, head = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    green, red = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFC7CE")
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"
    ws.column_dimensions["A"].width = 34
    for col in "BCDEFGHIJKLMNOPQRSTUV":
        ws.column_dimensions[col].width = 16
    ws.append([f"Final test on {TEST}: {', '.join(SYSTEMS[k] for k in data)}"])
    ws["A1"].font = Font(bold=True, size=13)
    for title, cols, rows in tables:
        ws.append([])
        ws.append([title])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.font, c.fill, c.alignment = bold, head, Alignment(wrap_text=True, vertical="top")
        for row in rows:
            ws.append([("n/a" if v != v else round(v, 3)) if isinstance(v, float) else v for v in row])
            for c in ws[ws.max_row]:
                if isinstance(c.value, str) and c.value.endswith("PASS"):
                    c.fill = green
                elif isinstance(c.value, str) and c.value.endswith("FAIL"):
                    c.fill = red

    # one row per chat: what each system extracted
    ws = wb.create_sheet("Per chat")
    ws.append(["id", "Chat text", "Expected PII"] + [SYSTEMS[k] for k in data])
    for c in ws[1]:
        c.font, c.fill = bold, head
    for col, w in zip("ABCDEF", (20, 60, 34, 34, 34, 34)):
        ws.column_dimensions[col].width = w
    by_id = {k: {r["id"]: r for r in res} for k, res in data.items()}
    for r in data["jev"]:
        row = [r["id"], r["text"], "\n".join(f"{g['label']}: {g['text']}" for g in r["labels"]["spans"])]
        for k in data:
            x = by_id[k][r["id"]]
            cell = f"FAILED: {x['error']}" if x.get("error") else "\n".join(f"{i['label']}: {i['text']}" for i in x["extracted"])
            if x.get("unsure"):
                cell += "\n[sent to a person: " + ", ".join(x["unsure"]) + "]"
            row.append(cell)
        ws.append(row)
        for c in ws[ws.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"

    path = RUNS / f"final_comparison_{dt.datetime.now():%Y%m%d-%H%M%S}.xlsx"
    wb.save(path)
    print(f"\nComparison: {path}")


# ---------------------------------------------------------------- 5. redacted file (opt-in)
def gold_redact(text, spans):
    """The expected redaction, from the labelled character spans."""
    for g in sorted(spans, key=lambda g: -g["start"]):
        text = text[:g["start"]] + f"[{g['label']}]" + text[g["end"]:]
    return text


def redact():
    """Redacts the test chats with what Jev + Haiku and Haiku alone extracted, next to the expected redaction.

    Over-redacted = an extracted item that isn't expected PII (text hidden for no reason).
    Missed = expected PII that is still visible.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    systems = ["jev", "haiku"]
    data = {k: {r["id"]: r for r in load(k)} for k in systems}
    stamp = f"{dt.datetime.now():%Y%m%d-%H%M%S}"
    out, totals = [], {k: Counter() for k in systems}
    for rid, base in sorted(data["jev"].items()):
        rec = {"id": rid, "original": base["text"], "expected": gold_redact(base["text"], base["labels"]["spans"])}
        for k in systems:
            r = data[k][rid]
            gh, ph, _ = P.span_match(r["text"], r["extracted"], r["labels"]["spans"])
            over = [i for i, h in zip(r["extracted"], ph) if not h]
            missed = [g for g, h in zip(r["labels"]["spans"], gh) if not h]
            rec[k] = {"redacted": P.redact_pii(r["text"], r["extracted"]), "redacted_items": r["extracted"],
                      "over_redacted": over, "missed": [{"label": g["label"], "text": g["text"]} for g in missed],
                      "sent_to_person": r.get("unsure", [])}
            t = totals[k]
            t["items redacted"] += len(r["extracted"])
            t["items over-redacted"] += len(over)
            t["chats over-redacted"] += bool(over)
            t["clean chats changed"] += bool(r["extracted"]) and not r["labels"]["has_pii"]
            t["expected PII missed"] += len(missed)
            t["chars hidden that weren't PII"] += sum(len(i["text"]) for i in over)
        out.append(rec)

    RUNS.mkdir(exist_ok=True)
    jpath = RUNS / f"redacted_test_{stamp}.json"
    jpath.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")

    bold, head = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    amber, red = PatternFill("solid", fgColor="FFEB9C"), PatternFill("solid", fgColor="FFC7CE")
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    ws.append([f"Redaction of {TEST} ({len(out)} chats, {sum(len(d['labels']['spans']) for d in data['jev'].values())} expected PII items)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["", SYSTEMS["jev"], SYSTEMS["haiku"]])
    for c in ws[3]:
        c.font, c.fill = bold, head
    for m in totals["jev"]:
        ws.append([m, totals["jev"][m], totals["haiku"][m]])
    ws.append(["redaction precision (items)"] + [round(1 - totals[k]["items over-redacted"] / totals[k]["items redacted"], 3)
                                                 for k in systems])
    ws.append([])
    ws.append(["Amber = text hidden that wasn't PII (over-redaction). Red = expected PII left visible."])

    ws = wb.create_sheet("Redacted chats")
    cols = ["id", "Original", "Expected", f"{SYSTEMS['jev']} redacted", f"{SYSTEMS['jev']}: hidden but not PII",
            f"{SYSTEMS['haiku']} redacted", f"{SYSTEMS['haiku']}: hidden but not PII"]
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill = bold, head
    for col, w in zip("ABCDEFG", (20, 50, 50, 50, 24, 50, 24)):
        ws.column_dimensions[col].width = w
    for rec in out:
        row = [rec["id"], rec["original"], rec["expected"]]
        for k in systems:
            x = rec[k]
            cell = x["redacted"]
            if x["sent_to_person"]:
                cell += "\n[sent to a person: " + ", ".join(x["sent_to_person"]) + "]"
            row += [cell, "\n".join(f"{i['label']}: {i['text']}" for i in x["over_redacted"])]
        ws.append(row)
        rr = ws.max_row
        for c in ws[rr]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for k, col in zip(systems, (4, 6)):
            if rec[k]["missed"]:
                ws.cell(rr, col).fill = red
            elif rec[k]["over_redacted"]:
                ws.cell(rr, col).fill = amber
                ws.cell(rr, col + 1).fill = amber
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:G{ws.max_row}"
    xpath = RUNS / f"redacted_test_{stamp}.xlsx"
    wb.save(xpath)

    print(f"{'':32}{SYSTEMS['jev']:>14}{SYSTEMS['haiku']:>14}")
    for m in totals["jev"]:
        print(f"{m:32}{totals['jev'][m]:>14}{totals['haiku'][m]:>14}")
    print(f"\nsaved {jpath}\nsaved {xpath}")


if __name__ == "__main__":
    if sys.argv[1] == "presidio-import":
        import_presidio(sys.argv[2])
    else:
        {"jev": run_jev, "haiku": run_haiku, "presidio": run_presidio, "colab": build_colab_notebook,
         "compare": compare, "redact": redact}[sys.argv[1]]()
