"""PII detection with Jev, and extraction with Claude Haiku 4.5, traced to LangSmith.

Two steps per chat:
    1. detect_pii(text)             Jev answers 7 yes/no questions, one per category
    2. extract_pii(text, detection) Claude Haiku 4.5 (via OpenRouter) copies out the exact
                                    PII, only for chats Jev flags or is unsure about

Categories: FULL_NAME, CARD_NUMBER, NRIC_PASSPORT, ADDRESS, PHONE, EMAIL, DOB

Only the chat text reaches the models. Any label columns in the CSV (has_pii, pii_types,
tags, spans_json) are used afterwards, by score(), to grade the answers.

Every chat becomes one LangSmith trace:
    pii_chat (chain)         final result, total latency, tokens and cost
      jev_detect (llm)       the Jev call
      llm_extract (llm)      the Haiku call (flagged chats only)
plus feedback scores (correct_has_pii, correct_<CAT>, span_recall, span_precision, correct_all).

Usage:
    python pii_detect.py ecom_realistic_dev_150.csv
    python pii_detect.py ecom_realistic_dev_150.csv --limit 5
    python pii_detect.py ecom_realistic_dev_150.csv --ids ECOM_REALISTIC-0006,ECOM_REALISTIC-0099
    python pii_detect.py ecom_realistic_dev_150.csv --no-trace

Keys are read from the environment, then .env, then the Windows user environment (never printed):
    TYPESAFE_API_KEY     Jev
    OPENROUTER_API_KEY   Claude Haiku 4.5
    LANGSMITH_API_KEY    tracing (LANGSMITH_PROJECT optional, default "pii-detection-with-jev")

The final test against Haiku alone and Presidio, and the redacted file, are in compare_test.py.
"""
import argparse
import csv
import datetime as dt
import json
import os
import random
import re
import sys
import threading
import time
import urllib.error
import urllib.request
import uuid
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

HERE = Path(__file__).parent

# ---------------------------------------------------------------- models and prices
JEV_ENDPOINT = "https://api.typesafe.ai/v1/systemone"
JEV_MODEL = "jev-latest"
JEV_PRICE_INPUT = 0.042  # USD per million input tokens, jev-1.13.0 (docs.typesafe.ai/models.md); output is free
JEV_PRICED_VERSION = "jev-1.13.0"

OPENROUTER_ENDPOINT = "https://openrouter.ai/api/v1/chat/completions"
LLM_MODEL = "anthropic/claude-haiku-4.5"
LLM_PRICES = (1.00, 5.00)  # input, output per million; fallback when OpenRouter doesn't report the cost

YES_AT = 0.5  # p_yes at or above this = yes

CATEGORIES = ["FULL_NAME", "CARD_NUMBER", "NRIC_PASSPORT", "ADDRESS", "PHONE", "EMAIL", "DOB"]

# risk if leaked: high-risk categories must never leak (caught by Jev, or flagged for a person)
RISK = {"CARD_NUMBER": "high", "NRIC_PASSPORT": "high",
        "FULL_NAME": "medium", "ADDRESS": "medium", "PHONE": "medium", "EMAIL": "medium", "DOB": "medium"}
# "% sure" = max(p, 1 - p). A category goes to human review when Jev is less sure than this.
REVIEW_BELOW = {"high": 0.95,    # p between 0.05 and 0.95
                "medium": 0.80}  # p between 0.20 and 0.80

# release gate
COST_CAP_PER_1K = 0.50  # USD per 1,000 chats, Jev + Haiku together
PRECISION_MIN = 0.90    # share of extracted items that are real PII (all categories); over-flagging hurts usefulness


def item_precision(rows):
    """(correct, extracted) over every extracted item in labelled rows, high and medium risk together."""
    hit = n = 0
    for r in rows:
        ph = span_match(r["text"], r["extracted"], r["labels"]["spans"])[1]
        hit, n = hit + sum(ph), n + len(ph)
    return hit, n


# ---------------------------------------------------------------- keys
def read_dotenv(path=HERE / ".env"):
    """KEY=value lines from a .env file next to this script (see .env.example)."""
    out = {}
    if path.exists():
        for line in path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1)
                out[k.strip()] = v.strip().strip('"').strip("'")
    return out


DOTENV = read_dotenv()


def env_key(name):
    """The environment first, then .env, then (on Windows) the user environment in the registry."""
    val = os.environ.get(name) or DOTENV.get(name)
    if not val and sys.platform == "win32":
        import winreg
        try:
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, "Environment") as k:
                val, _ = winreg.QueryValueEx(k, name)
        except OSError:
            val = None
    if val:
        os.environ[name] = val  # so SDKs that read os.environ see it
    return val


# ---------------------------------------------------------------- questions (Jev)
COUNTS_ALSO = ("Count it even when it is written in another language or script, spelled out in words "
               "(\"nine zero zero eight\"), disguised (\"at\" / \"dot\" / \"[at]\"), split across several "
               "messages, or given and then corrected.")

QUESTIONS = {
    "FULL_NAME": {
        "type": "noul",
        "instructions": (
            "Does `chat` contain the full name of a private person, meaning a given name together with a "
            "family name (for example 'Bill Clark', 'Ms Pei Shan Lee', 'Siti Nurhaliza binte Ahmad', or a "
            "full name written in Tamil or Chinese)? A first name on its own ('Hi Rachel') does NOT count. "
            "Product names that contain name-like words ('Jack Leather Wallet', 'Grace Linen Dress', "
            "'Mark II Headphones'), place names and company names do NOT count. A name that appears only "
            "as part of an email address ('rachel dot chua at outlook dot com', 'nurul.nair@hotmail.com') "
            "does NOT count here, because that is an email. " + COUNTS_ALSO),
    },
    "CARD_NUMBER": {
        "type": "noul",
        "instructions": (
            "Does `chat` contain a full payment card number (credit or debit card, usually 15 or 16 digits, "
            "with any spacing or dashes, or read out in words) that a customer gives as their card? "
            "The last 4 digits of a card do NOT count. Order numbers, tracking numbers, SKUs, coupon codes "
            "and amounts do NOT count, even when they are long, so use what the chat says the number is. "
            + COUNTS_ALSO),
    },
    "NRIC_PASSPORT": {
        "type": "noul",
        "instructions": (
            "Does `chat` contain a national identity or passport number of a private person, such as a "
            "Singapore NRIC or FIN (a letter S, T, F, G or M, then 7 digits, then a letter, e.g. "
            "'S1234567D') or a passport number? Order numbers, tracking numbers and coupon codes do NOT "
            "count. " + COUNTS_ALSO),
    },
    "ADDRESS": {
        "type": "noul",
        "instructions": (
            "Does `chat` contain a specific home, office or delivery address of a private person, such as a "
            "block or street number with a street name, a unit number like '#09-68', or a postal code "
            "(e.g. '337 Yishun Ring Road, #29-18, Singapore 365868')? A city, town or area name on its own "
            "('Tampines', 'Kuala Lumpur'), a store or outlet location ('our Yishun outlet') and the "
            "company's own address do NOT count. " + COUNTS_ALSO),
    },
    "PHONE": {
        "type": "noul",
        "instructions": (
            "Does `chat` contain the phone number of a private person (e.g. '+65 9777 5204', '86690967', "
            "or pieces of one number given across several messages)? The company's hotline or customer "
            "service number does NOT count. Order numbers, tracking numbers, SKUs and amounts do NOT "
            "count, so use what the chat says the number is. " + COUNTS_ALSO),
    },
    "EMAIL": {
        "type": "noul",
        "instructions": (
            "Does `chat` contain the email address of a private person (e.g. 'klim84@icloud.com', or "
            "'rachel dot chua at outlook dot com')? The company's own addresses (such as "
            "returns@shopsite.com or a support address) do NOT count. " + COUNTS_ALSO),
    },
    "DOB": {
        "type": "noul",
        "instructions": (
            "Does `chat` contain a private person's date of birth (e.g. 'DOB 08/09/2000', 'born on 3 May "
            "1990', or a birth date in another language)? Delivery dates, restock dates, order dates and "
            "a person's age do NOT count. " + COUNTS_ALSO),
    },
}
QUESTIONS_V1 = QUESTIONS  # first wording: one long paragraph per category


# v2: short positive question + true/false criteria with examples (docs.typesafe.ai/primitives/advanced.md).
# Examples are invented (not copied from the dev chats) to limit overfitting.
FOCUS = ("Read every message. A value split across several messages, spelled out in words, written in "
         "another language, or given and then corrected still counts.")


NAME_FOCUS = ("Judge only words used as a person's name, such as when someone introduces themselves or gives the "
              "name on an order. Ignore email addresses and product names completely.")


def _q(question, true_what, true_examples, false_what, false_examples, focus=FOCUS):
    return {
        "type": "noul",
        "instructions": {"question": question, "focus": focus},
        "criteria": {
            "true": {"what": true_what, "examples": true_examples},
            "false": {"what": false_what, "examples": false_examples},
        },
    }


QUESTIONS_V2 = {
    "FULL_NAME": _q(
        "Does `chat` name a private person with their full name?",
        "Someone's given name and family name are written together as their name, in any language or script.",
        ["The name on the order is Daniel Koh.", "This is Mrs Farah Ismail speaking.", "Recipient: Tan Wei Ling",
         "என் பெயர் பிரியா ராமன்.", "我叫陈美玲。"],
        "The chat uses only a single first name for people, or the only name-like words are inside an email "
        "address or a product name.",
        ["Agent: Hi Jason, welcome back! (no other name given)", "Thanks, Mei!",
         "Please update my email to daniel dot koh at gmail dot com.", "Send it to mei.ong@hotmail.com",
         "Is the Oliver Desk Lamp in stock?", "Does the Daisy Cotton Shirt come in blue?"],
        focus=NAME_FOCUS),
    "CARD_NUMBER": _q(
        "Does `chat` contain a customer's full payment card number?",
        "The full number of a credit or debit card: 15 or 16 digits, written with spaces, dashes, no gaps, or spelled out in words.",
        ["My card is 5412 7534 9821 0043.", "Charge it to 3714-496353-98431.",
         "card number four five three nine one four eight eight zero three four three six four six seven"],
        "Only the last few digits of a card, or a long number the chat calls a tracking, order, receipt or reference number.",
        ["the card ending 4417", "Tracking number 7712 0938 4456 2201 hasn't moved.",
         "Where is the refund for SG4419027735?", "Use my saved Visa."]),
    "NRIC_PASSPORT": _q(
        "Does `chat` contain a person's national ID or passport number?",
        "A government-issued number that identifies a person, given as their ID: a national identity card number "
        "from any country, or a passport number.",
        ["My ID card number is S8812345B.", "Passport number K1234567N", "Customs asked for my ID: 880512-14-5523"],
        "Order, receipt, tracking, account and coupon codes, even when they mix letters and digits.",
        ["I need a receipt for SG7720194463.", "Order S00128834", "Coupon SAVE15 failed."]),
    "ADDRESS": _q(
        "Does `chat` contain a private person's home, office or delivery address?",
        "A specific place: a street number and street name, a block, a unit number like #05-12, or a 6-digit postal code.",
        ["Deliver to Blk 221 Bedok North Ave 3 #08-115 Singapore 460221.",
         "My office is at 18 Cross Street, #12-01, Singapore 048423."],
        "Only a town, city or area name, or the company's own store or outlet.",
        ["Do you ship to Punggol?", "Can I collect it at your Orchard store?", "How long does delivery to Penang take?"]),
    "PHONE": _q(
        "Does `chat` contain a private person's phone number?",
        "A number the customer gives so they or someone else can be contacted. It may be split across messages or spelled out in words.",
        ["Call me on 9123 4567.", "+65 8765 4321", "Customer: It's nine one two three... / Customer: four five six seven.",
         "Agent: What's your contact number? / Customer: 8234 1199."],
        "The company's own hotline or service numbers, a number the customer asks about as the company's number, "
        "and order, tracking or amount numbers.",
        ["Our hotline is 6123 4567.", "Is 1800 111 2222 your returns line?", "Is 6345 1234 the right number for returns?",
         "I was charged twice for order 88123456."]),
    "EMAIL": _q(
        "Does `chat` contain a private person's email address?",
        "An email address belonging to a customer or another private person, written normally or with 'at' and 'dot' spelled out.",
        ["jason.lim88@gmail.com", "mei dot ong at hotmail dot com"],
        "The company's own email addresses.",
        ["Write to support@shopsite.com.", "Send it to returns@shopsite.com."]),
    "DOB": _q(
        "Does `chat` contain a person's date of birth?",
        "A date the chat gives as someone's birth date, in any format or language.",
        ["DOB 14/02/1995", "I was born on 3 March 1988.", "我的生日是1990年5月6日。"],
        "Delivery, order, expiry or restock dates, and ages.",
        ["It will arrive on 12 June.", "The code expired on 01/04/2026.", "I'm 34 years old."]),
}

QUESTION_SETS = {"v1": QUESTIONS_V1, "v2": QUESTIONS_V2}
QUESTIONS = QUESTIONS_V2  # default

# Jev is weak at reading spelled-out digits, so for v2 runs of 3+ digit words become numerals before
# Jev sees the chat ("nine zero zero eight" -> "9008"). Haiku always gets the original text.
DIGIT_WORDS = {"zero": "0", "oh": "0", "o": "0", "one": "1", "two": "2", "three": "3", "four": "4",
               "five": "5", "six": "6", "seven": "7", "eight": "8", "nine": "9"}
_DIGIT_RUN = re.compile(r"\b(?:%s)(?:[\s,-]+(?:%s)){2,}\b" % (("|".join(DIGIT_WORDS),) * 2), re.IGNORECASE)


def normalize_digits(text):
    return _DIGIT_RUN.sub(lambda m: "".join(DIGIT_WORDS[w.lower()] for w in re.findall(r"[A-Za-z]+", m.group())),
                          text)


# ---------------------------------------------------------------- HTTP
def post_json(url, payload, key, label, retries=5, timeout=60):
    """POST JSON with a bearer key; retries rate limits, server errors and network errors."""
    body = json.dumps(payload).encode()
    for attempt in range(retries + 1):
        req = urllib.request.Request(
            url, data=body, method="POST",
            headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"},
        )
        try:
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                return json.loads(resp.read())
        except urllib.error.HTTPError as e:
            if e.code in (408, 429, 500, 502, 503, 529) and attempt < retries:
                time.sleep(min(2 ** attempt + random.random(), 30))
                continue
            raise RuntimeError(f"{label} HTTP {e.code} {e.reason}: {e.read().decode(errors='replace')[:500]}") from None
        except urllib.error.URLError:
            if attempt < retries:
                time.sleep(min(2 ** attempt + random.random(), 30))
                continue
            raise


# ---------------------------------------------------------------- step 1: detect (Jev)
def detect_pii(text, key=None, questions=None):
    """One Jev call: 7 yes/no questions, one per category.

    Each answer is Jev's probability of "yes" (p). The answer is yes when p >= 0.5, and
    confidence is how sure Jev is of that answer: max(p, 1 - p), so 0.5 = coin flip, 1.0 = certain.
    A category goes to human review ("unsure") when confidence is below its risk tier's bar.
    """
    key = key or env_key("TYPESAFE_API_KEY")
    if not key:
        raise RuntimeError("TYPESAFE_API_KEY is not set.")
    start = time.perf_counter()
    data = post_json(JEV_ENDPOINT, {"model": JEV_MODEL, "state": {"chat": text}, "questions": questions or QUESTIONS}, key, "Jev")
    ms = (time.perf_counter() - start) * 1000

    cats = {}
    for c in CATEGORIES:
        p = float(data["answers"][c]["noul"])
        cats[c] = {"p_yes": p, "answer": "yes" if p >= YES_AT else "no", "confidence": max(p, 1 - p)}
    found = [c for c in CATEGORIES if cats[c]["answer"] == "yes"]
    # sent to a person when Jev is less sure than the category's risk tier allows
    review = [c for c in CATEGORIES if cats[c]["confidence"] < REVIEW_BELOW[RISK[c]]]
    tokens = data["usage"]["input_tokens"]
    return {
        "has_pii": bool(found), "categories": cats, "types_found": found, "unsure": review,
        "input_tokens": tokens, "cost": tokens * JEV_PRICE_INPUT / 1e6, "latency_ms": ms,
        "model_version": data.get("model"), "raw_answers": data["answers"],
    }


# ---------------------------------------------------------------- step 2: extract (Claude Haiku 4.5)
LLM_SYSTEM = """You extract personally identifiable information (PII) from an e-commerce customer support chat.

Return every piece of PII about a private person (the customer, or someone they mention such as a gift recipient) in these categories:
- FULL_NAME: a given name together with a family name, in any language or script, without a title (return "Pei Shan Lee", not "Ms Pei Shan Lee").
- CARD_NUMBER: a full payment card number.
- NRIC_PASSPORT: a national identity number from any country, or a passport number.
- ADDRESS: a specific home, office or delivery address, including the unit number and postal code.
- PHONE: a private person's phone number.
- EMAIL: a private person's email address.
- DOB: a date of birth.

Rules:
- Each "text" must be copied EXACTLY, character for character, from the chat. Do not fix spelling, spacing or punctuation, do not join pieces, and do not add or drop words. Leave out surrounding words and the trailing full stop.
- If one item is split across several messages (e.g. a phone number given in two parts), return each piece as its own item with the same label.
- If a customer gives a value and then corrects it, return both values.
- Spoken or disguised forms count: return them as written ("nine zero zero eight", "rachel dot chua at outlook dot com").
- Do NOT return: the company's hotline, store or support emails, order / tracking / SKU / coupon numbers, the last 4 digits of a card, a first name on its own, product names, city or area names on their own, store locations, delivery or restock dates, or ages.
- The hint lists the categories a screening model flagged. Check them, but return whatever is actually in the chat, and return an empty list if there is no PII."""

LLM_SCHEMA = {
    "type": "object",
    "properties": {
        "pii": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {"label": {"type": "string", "enum": CATEGORIES}, "text": {"type": "string"}},
                "required": ["label", "text"],
                "additionalProperties": False,
            },
        }
    },
    "required": ["pii"],
    "additionalProperties": False,
}


def extract_pii(text, detection, key=None):
    """One Haiku call through OpenRouter. Keeps only items whose text appears verbatim in the chat.

    Returns {"items": [{label, text}], "rejected": [{label, text}], usage and cost}.
    """
    key = key or env_key("OPENROUTER_API_KEY")
    if not key:
        raise RuntimeError("OPENROUTER_API_KEY is not set.")
    hint = sorted(set(detection["types_found"]) | set(detection["unsure"]))
    user = json.dumps({"chat": text, "hint_categories": hint}, ensure_ascii=False)
    start = time.perf_counter()
    data = post_json(OPENROUTER_ENDPOINT, {
        "model": LLM_MODEL,
        "max_tokens": 1024,
        "temperature": 0,
        "messages": [{"role": "system", "content": LLM_SYSTEM}, {"role": "user", "content": user}],
        "response_format": {"type": "json_schema", "json_schema": {"name": "pii", "strict": True, "schema": LLM_SCHEMA}},
        "provider": {"require_parameters": True},  # only route to providers that honour the schema
        "usage": {"include": True},                # report the actual cost
    }, key, "OpenRouter", timeout=120)
    ms = (time.perf_counter() - start) * 1000
    if "error" in data:
        raise RuntimeError(f"OpenRouter error: {data['error']}")
    choice = data["choices"][0]
    if choice.get("finish_reason") not in ("stop", "end_turn", None):
        raise RuntimeError(f"OpenRouter finish_reason={choice.get('finish_reason')}")
    answer = json.loads(choice["message"].get("content") or "{}")

    items, rejected, seen = [], [], set()
    for it in answer.get("pii", []):
        t = it.get("text", "").strip()
        if it.get("label") in CATEGORIES and t and t in text:
            if (it["label"], t) not in seen:
                seen.add((it["label"], t))
                items.append({"label": it["label"], "text": t})
        else:
            rejected.append(it)

    u = data.get("usage") or {}
    tin, tout = u.get("prompt_tokens", 0), u.get("completion_tokens", 0)
    cost = u.get("cost")
    if cost is None:
        cost = (tin * LLM_PRICES[0] + tout * LLM_PRICES[1]) / 1e6
    return {"items": items, "rejected": rejected, "input_tokens": tin, "output_tokens": tout,
            "cost": float(cost), "latency_ms": ms,
            "served_by": f"{data.get('model')} via {data.get('provider', 'openrouter')}", "raw": answer}


# ---------------------------------------------------------------- grading helpers
def read_labels(row):
    """Ground truth from the CSV row, or None when the file has no label columns."""
    if not row.get("has_pii"):
        return None
    types = {t for t in (row.get("pii_types") or "").split("|") if t}
    spans = json.loads(row.get("spans_json") or "[]")
    return {"has_pii": row["has_pii"].strip().upper() == "TRUE", "types": types, "spans": spans}


def char_ranges(text, needle):
    """Every (start, end) where needle occurs in text."""
    out, i = [], text.find(needle)
    while needle and i != -1:
        out.append((i, i + len(needle)))
        i = text.find(needle, i + 1)
    return out


def span_match(text, items, gold_spans):
    """Overlap matching of extracted items against gold spans.

    Returns (gold_hit[list of bool], pred_hit[list of bool], label_ok[list of bool for matched pairs]).
    """
    pred_ranges = [char_ranges(text, it["text"]) for it in items]
    gold_hit, label_ok = [], []
    for g in gold_spans:
        hit = False
        for it, rngs in zip(items, pred_ranges):
            if any(s < g["end"] and g["start"] < e for s, e in rngs):
                hit = True
                label_ok.append(it["label"] == g["label"])
                break
        gold_hit.append(hit)
    pred_hit = [any(s < g["end"] and g["start"] < e for s, e in rngs for g in gold_spans) for rngs in pred_ranges]
    return gold_hit, pred_hit, label_ok


# ---------------------------------------------------------------- one chat (traced)
def make_row_fn(cfg):
    from langsmith import traceable
    from langsmith.run_helpers import get_current_run_tree

    def set_run(usage, meta):
        run = get_current_run_tree()
        if run is not None:
            run.set(usage_metadata=usage, metadata=meta)

    @traceable(name="jev_detect", run_type="llm",
               metadata={"ls_provider": "typesafe", "ls_model_name": JEV_MODEL})
    def jev_step(state):
        chat = normalize_digits(state["chat"]) if cfg.questions != "v1" else state["chat"]
        d = detect_pii(chat, cfg.typesafe_key, QUESTION_SETS[cfg.questions])
        set_run({"input_tokens": d["input_tokens"], "output_tokens": 0, "total_tokens": d["input_tokens"],
                 "input_cost": d["cost"], "output_cost": 0.0, "total_cost": d["cost"]},
                {"jev_model_version": d["model_version"], "round_trip_ms": round(d["latency_ms"])})
        return d

    @traceable(name="llm_extract", run_type="llm",
               metadata={"ls_provider": "openrouter", "ls_model_name": LLM_MODEL})
    def llm_step(state, detection):
        x = extract_pii(state["chat"], detection, cfg.openrouter_key)
        set_run({"input_tokens": x["input_tokens"], "output_tokens": x["output_tokens"],
                 "total_tokens": x["input_tokens"] + x["output_tokens"], "total_cost": x["cost"]},
                {"served_by": x["served_by"], "round_trip_ms": round(x["latency_ms"])})
        return x

    # labels never reach the models, and are left out of the trace inputs
    @traceable(name="pii_chat", run_type="chain", process_inputs=lambda d: dict(d["inputs"]))
    def pii_chat(inputs, labels):
        start = time.perf_counter()
        det = jev_step(inputs)
        ext = None
        if det["types_found"] or det["unsure"]:
            ext = llm_step(inputs, {k: det[k] for k in ("types_found", "unsure")})
        items = ext["items"] if ext else []
        got = {it["label"] for it in items}
        disagree = sorted(set(det["types_found"]) ^ got) if ext else []

        res = {
            "has_pii": det["has_pii"],
            "categories": det["categories"], "types_found": det["types_found"], "unsure": det["unsure"],
            "extracted": items, "rejected": ext["rejected"] if ext else [], "disagree": disagree,
            "jev_model_version": det["model_version"],
            "jev_tokens": det["input_tokens"], "jev_cost": det["cost"], "jev_ms": det["latency_ms"],
            "llm_called": ext is not None,
            "llm_tokens_in": ext["input_tokens"] if ext else 0, "llm_tokens_out": ext["output_tokens"] if ext else 0,
            "llm_cost": ext["cost"] if ext else 0.0, "llm_ms": ext["latency_ms"] if ext else 0.0,
            "error": "",
        }
        res["cost_usd"] = res["jev_cost"] + res["llm_cost"]
        if labels:
            res["correct"] = {"has_pii": det["has_pii"] == labels["has_pii"]}
            for c in CATEGORIES:
                res["correct"][c] = (c in det["types_found"]) == (c in labels["types"])
            gold_hit, pred_hit, label_ok = span_match(inputs["chat"], items, labels["spans"])
            res["span_recall"] = sum(gold_hit) / len(gold_hit) if gold_hit else None
            res["span_precision"] = sum(pred_hit) / len(pred_hit) if pred_hit else None
            res["correct_all"] = all(res["correct"].values()) and all(gold_hit) and all(pred_hit) and all(label_ok)
        res["total_ms"] = (time.perf_counter() - start) * 1000
        return res

    def run_row(row, source):
        inputs = {"chat": row["text"]}
        labels = read_labels(row)
        run_id = uuid.uuid4()
        tags = [t for t in (row.get("tags") or "").split("|") if t]
        extra = {
            "run_id": run_id,
            "metadata": {"id": row["id"], "source_file": source, "batch_id": cfg.batch_id,
                         "jev_model": JEV_MODEL, "llm_model": LLM_MODEL, "questions": cfg.questions},
            "tags": [source, f"batch:{cfg.batch_id}", f"questions:{cfg.questions}"] + tags,
        }
        res = pii_chat(inputs, labels, langsmith_extra=extra)
        if cfg.ls_client and labels:
            try:
                fb = {f"correct_{k}": int(v) for k, v in res["correct"].items()}
                fb["correct_all"] = int(res["correct_all"])
                for k in ("span_recall", "span_precision"):
                    if res[k] is not None:
                        fb[k] = res[k]
                for k, v in fb.items():
                    cfg.ls_client.create_feedback(run_id=run_id, trace_id=run_id, session_id=cfg.ls_project_id,
                                                  key=k, score=v)
            except Exception as e:  # feedback is nice-to-have; never fail the row on it
                with cfg.print_lock:
                    print(f"  (feedback for {row['id']} not sent: {e})")
        res.update(id=row["id"], text=row["text"], source_file=source, labels=labels, tags=tags,
                   trace_id=str(run_id) if cfg.ls_client else "")
        return res

    return run_row


# ---------------------------------------------------------------- batch
def run(csv_path, cfg, limit=None, ids=None):
    path = Path(csv_path) if Path(csv_path).is_absolute() else HERE / csv_path
    with open(path, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    if not rows or "text" not in rows[0] or "id" not in rows[0]:
        sys.exit(f"{path.name} needs 'id' and 'text' columns.")
    if ids:
        rows = [r for r in rows if r["id"] in ids]
    rows = rows[:limit]

    tracing = "off" if not cfg.ls_client else f"LangSmith project '{os.environ['LANGSMITH_PROJECT']}'"
    print(f"batch {cfg.batch_id}: {len(rows)} chats | Jev (questions {cfg.questions}) + {LLM_MODEL} | tracing: {tracing}")
    run_row = make_row_fn(cfg)
    results, t0 = [], time.perf_counter()
    with ThreadPoolExecutor(max_workers=cfg.workers) as pool:
        futs = {pool.submit(run_row, r, path.name): r for r in rows}
        for i, fut in enumerate(as_completed(futs), 1):
            r = futs[fut]
            try:
                results.append(fut.result())
            except Exception as e:  # kept in the results, never dropped
                results.append({"id": r["id"], "text": r["text"], "source_file": path.name,
                                "labels": read_labels(r), "error": str(e)[:500]})
                with cfg.print_lock:
                    print(f"  FAILED {r['id']}: {e}")
            if i % 25 == 0 or i == len(rows):
                with cfg.print_lock:
                    print(f"  {i}/{len(rows)} done")
    wall = time.perf_counter() - t0
    if cfg.ls_client:
        cfg.ls_client.flush()
    results.sort(key=lambda r: r["id"])
    summary = score(results, wall)
    write_results(cfg, results, summary)
    if cfg.ls_client:
        print(f"\nTraces: LangSmith project '{os.environ['LANGSMITH_PROJECT']}', filter by tag batch:{cfg.batch_id}")
    return results


def write_results(cfg, results, summary):
    """One Excel workbook per run.

    Results: one row per chat. For the overall question and each of the 7 categories, the
             expected answer sits next to Jev's answer and how sure Jev was ("Yes 98%").
             Green = Jev matches the expected answer, red = it doesn't.
    Summary: precision / recall / F1, per category scores, confidence check, extraction, cost.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    out_dir = HERE / "runs"
    out_dir.mkdir(exist_ok=True)
    path = out_dir / f"pii_results_{cfg.batch_id}_{cfg.questions}.xlsx"

    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    top = Alignment(vertical="top", wrap_text=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    group_fills = [PatternFill("solid", fgColor=c) for c in ("BDD7EE", "D9E1F2")]
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    amber = PatternFill("solid", fgColor="FFEB9C")
    edge = Side(style="medium", color="808080")
    yn = lambda b: "Yes" if b else "No"

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    # ---- two header rows: group names (merged) over Expected | Jev
    nice = {"FULL_NAME": "Full name", "CARD_NUMBER": "Card number", "NRIC_PASSPORT": "NRIC / passport",
            "ADDRESS": "Address", "PHONE": "Phone", "EMAIL": "Email", "DOB": "Date of birth"}
    groups = [(f"{nice[c]}\n({RISK[c]} risk)", c) for c in CATEGORIES]
    top_row = ["id", "Chat text"]
    sub_row = ["", ""]
    for name, _ in groups:
        top_row += [name, ""]
        sub_row += ["Expected", "Jev (% sure)"]
    top_row += ["All correct?", "Extracted by Haiku", "Review note"]
    sub_row += ["", "", ""]
    ws.append(top_row)
    ws.append(sub_row)

    n_cols = len(top_row)
    first_group_col = 3
    for col in range(1, n_cols + 1):
        for row in (1, 2):
            cell = ws.cell(row, col)
            cell.font, cell.alignment, cell.fill = bold, center, head_fill
    for col in (1, 2, n_cols - 2, n_cols - 1, n_cols):  # single columns span both header rows
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
    for g in range(len(groups)):
        c0 = first_group_col + 2 * g
        ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + 1)
        for col in (c0, c0 + 1):
            for row in (1, 2):
                ws.cell(row, col).fill = group_fills[g % 2]

    widths = [20, 60] + [10, 12] * len(groups) + [10, 38, 34]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ---- one row per chat
    for r in results:
        lab = r.get("labels") or {}
        row = [r["id"], r["text"]]
        if r.get("error"):
            row += [""] * (2 * len(groups)) + ["", "", f"FAILED: {r['error']}"]
            ws.append(row)
            for c in ws[ws.max_row]:
                c.alignment, c.fill = top, red
            continue

        exp_types = lab.get("types", set())
        cells = []
        for c in CATEGORIES:
            cat = r["categories"][c]
            cells.append((yn(c in exp_types) if lab else "", cat["answer"] == "yes", c in r["unsure"],
                          (c in exp_types) if lab else None))
            row += [cells[-1][0], f"{yn(cat['answer'] == 'yes')} {cat['confidence']:.0%}"]

        note = []
        if r["unsure"]:
            note.append("Send to a person, Jev not sure enough about: " + ", ".join(r["unsure"]))
        if r["disagree"]:
            note.append("Jev and Haiku disagree on: " + ", ".join(r["disagree"]))
        if r["rejected"]:
            note.append(f"Haiku text not found in chat: {r['rejected']}")
        correct = r.get("correct_all", "")
        row += ["" if correct == "" else ("✓" if correct else "✗"),
                "\n".join(f"{i['label']}: {i['text']}" for i in r["extracted"]),
                "\n".join(note)]
        ws.append(row)

        rr = ws.max_row
        for c in ws[rr]:
            c.alignment = top
        for g, (exp_text, jev_yes, flagged, expected) in enumerate(cells):
            jcell = ws.cell(rr, first_group_col + 2 * g + 1)
            ecell = ws.cell(rr, first_group_col + 2 * g)
            ecell.alignment = jcell.alignment = Alignment(horizontal="center", vertical="top")
            if expected is not None:
                jcell.fill = green if jev_yes == expected else red
            if flagged:
                jcell.font = Font(bold=True, italic=True)
            if expected:  # make the expected "Yes" cells stand out
                ecell.font = bold
        ok_cell = ws.cell(rr, n_cols - 2)
        ok_cell.alignment = Alignment(horizontal="center", vertical="top")
        if correct != "":
            ok_cell.fill = green if correct else red
        if note:
            ws.cell(rr, n_cols).fill = amber

    # vertical lines between the groups
    for g in range(len(groups) + 1):
        col = first_group_col + 2 * g
        for row in range(1, ws.max_row + 1):
            cell = ws.cell(row, col)
            cell.border = Border(left=edge)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{get_column_letter(n_cols)}{ws.max_row}"

    # ---- Summary
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 30
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 13
    ws.append([f"Run {cfg.batch_id}", f"Jev ({JEV_MODEL}, questions {cfg.questions}) + {LLM_MODEL}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["How to read the Results sheet: 'Jev (% sure)' shows Jev's answer and how sure it was "
               "(50% = a guess, 100% = certain). Green = matches the expected answer, red = doesn't. "
               f"Bold italic = sent to a person (high risk: under {REVIEW_BELOW['high']:.0%} sure; "
               f"medium risk: under {REVIEW_BELOW['medium']:.0%} sure)."])
    for title, cols_, rows in summary:
        ws.append([])
        ws.append([title])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
        if cols_:
            ws.append(cols_)
            for c in ws[ws.max_row]:
                c.font, c.fill = bold, head_fill
        for row in rows:
            ws.append([("n/a" if v != v else round(v, 3)) if isinstance(v, float) else v for v in row])

    wb.save(path)
    print(f"\nResults: {path}")
    return path


# ---------------------------------------------------------------- scoring
def prf(tp, fp, fn):
    p = tp / (tp + fp) if tp + fp else float("nan")
    r = tp / (tp + fn) if tp + fn else float("nan")
    f = 2 * p * r / (p + r) if p + r and p == p and r == r else float("nan")
    return p, r, f


def release_gate(labelled, ok):
    """Release gate tables. The run passes only when all four checks pass:

    1. High-risk recall: zero leaks of card or NRIC/passport numbers. A piece of PII is SAFE when
       Jev said yes to its category and Haiku extracted it, or when the chat was sent to a person
       for that category (Jev under REVIEW_BELOW sure). Anything else is a LEAK.
       Medium-risk leaks are reported, but have no target.
    2. Zero invented text: every item Haiku returns must appear word-for-word in the chat.
    3. Cost: Jev + Haiku under COST_CAP_PER_1K per 1,000 chats. Speed is recorded, not gated.
    4. Precision: more than PRECISION_MIN of the extracted items (all categories) are real PII,
       so the output isn't over-redacted.
    """
    def gold_spans(r, c):
        return [g for g in r["labels"]["spans"] if g["label"] == c]

    rows, verdicts = [], {}
    for c in CATEGORIES:
        gold = [r for r in labelled if c in r["labels"]["types"]]
        caught = [r for r in gold if c in r["types_found"]]
        flagged = [r for r in gold if c not in r["types_found"] and c in r["unsure"]]
        missed = [r for r in gold if c not in r["types_found"] and c not in r["unsure"]]
        # detected, but the text itself wasn't extracted (so it couldn't be redacted) and nobody was asked
        not_extracted = []
        for r in caught:
            if c in r["unsure"]:
                continue
            hits = span_match(r["text"], r["extracted"], gold_spans(r, c))[0]
            if not all(hits):
                not_extracted.append(r)
        leaks = missed + not_extracted
        if RISK[c] == "high":
            verdict = "PASS" if not leaks else "FAIL"
            verdicts[c] = verdict
        else:
            verdict = "report only"
        rows.append([c, RISK[c], len(gold), len(caught), len(flagged), len(missed), len(not_extracted), len(leaks),
                     ", ".join(r["id"] for r in leaks) or "-", verdict])

    invented = sum(len(r["rejected"]) for r in ok)
    cost_1k = sum(r["cost_usd"] for r in ok) / len(ok) * 1000
    hit, n = item_precision(labelled)
    precision = hit / n if n else float("nan")
    checks = [
        ["1. High-risk recall (card, NRIC/passport)", "0 leaks",
         f"{sum(row[7] for row in rows if row[1] == 'high')} leaks",
         "PASS" if all(v == "PASS" for v in verdicts.values()) else "FAIL"],
        ["2. Invented text (Haiku items not in the chat)", "0", str(invented), "PASS" if invented == 0 else "FAIL"],
        ["3. Cost per 1,000 chats (Jev + Haiku)", f"under ${COST_CAP_PER_1K:.2f}", f"${cost_1k:.3f}",
         "PASS" if cost_1k < COST_CAP_PER_1K else "FAIL"],
        ["4. Precision (extracted items that are real PII, all categories)", f"over {PRECISION_MIN:.0%}",
         f"{precision:.1%} ({hit}/{n})", "PASS" if precision > PRECISION_MIN else "FAIL"],
    ]
    overall = "PASS" if all(c[3] == "PASS" for c in checks) else "FAIL"
    tables = [(f"RELEASE GATE: {overall}", ["Check", "Target", "This run", "Result"], checks),
              (f"Leaks by category (sent to a person when Jev is under {REVIEW_BELOW['high']:.0%} sure for high "
               f"risk, under {REVIEW_BELOW['medium']:.0%} for medium risk)",
               ["Category", "Risk", "Chats with it", "Caught by Jev", "Missed but sent to person",
                "Missed, not flagged", "Caught, not extracted", "LEAKS", "Leaked chats", "Gate"], rows)]

    # human workload
    sent = [r for r in labelled if r["unsure"]]
    sent_clean = [r for r in sent if not r["labels"]["has_pii"]]
    sent_high = [r for r in labelled if any(RISK[c] == "high" for c in r["unsure"])]
    tables.append(("Human review workload",
                   ["Chats sent to a person", "% of all chats", "of which clean", "sent for a high-risk category"],
                   [[len(sent), len(sent) / len(labelled), len(sent_clean), len(sent_high)]]))

    # what each tier's review bar would do at other settings
    for tier, bars in (("high", (0.65, 0.80, 0.90, 0.95, 0.99)), ("medium", (0.65, 0.70, 0.80, 0.90, 0.95))):
        sweep = []
        for bar in bars:
            misses = sends = 0
            for r in labelled:
                send = False
                for c in CATEGORIES:
                    if RISK[c] != tier:
                        continue
                    cat = r["categories"][c]
                    low = cat["confidence"] < bar
                    send = send or low
                    if c in r["labels"]["types"] and cat["answer"] == "no" and not low:
                        misses += 1
                sends += send
            sweep.append([f"under {bar:.0%} sure" + ("  (current)" if bar == REVIEW_BELOW[tier] else ""),
                          misses, sends, sends / len(labelled)])
        tables.append((f"{tier.title()}-risk review bar: what other settings would do (Jev detection only)",
                       ["Send to a person when", f"{tier.title()}-risk misses not flagged",
                        "Chats sent to a person", "% of all chats"], sweep))
    return tables


def score(results, wall=None):
    """Prints the scores and returns them as [(title, column names, rows)] for the Summary sheet."""
    ok = [r for r in results if not r.get("error")]
    failed = len(results) - len(ok)
    labelled = [r for r in ok if r.get("labels")]
    nan = float("nan")
    tables = []

    if labelled:
        tp = sum(r["has_pii"] and r["labels"]["has_pii"] for r in labelled)
        fp = sum(r["has_pii"] and not r["labels"]["has_pii"] for r in labelled)
        fn = sum(not r["has_pii"] and r["labels"]["has_pii"] for r in labelled)
        tn = len(labelled) - tp - fp - fn
        p, rc, f = prf(tp, fp, fn)
        tables.extend(release_gate(labelled, ok))
        tables.append(("Has PII? (per chat: any category Jev said yes to)",
                       ["Precision", "Recall", "F1", "Accuracy", "TP", "FP", "FN", "TN"],
                       [[p, rc, f, (tp + tn) / len(labelled), tp, fp, fn, tn]]))

        rows = []
        for c in CATEGORIES:
            ctp = sum(c in r["types_found"] and c in r["labels"]["types"] for r in labelled)
            cfp = sum(c in r["types_found"] and c not in r["labels"]["types"] for r in labelled)
            cfn = sum(c not in r["types_found"] and c in r["labels"]["types"] for r in labelled)
            cp, cr, cf = prf(ctp, cfp, cfn)
            conf = sum(r["categories"][c]["confidence"] for r in labelled) / len(labelled)
            rows.append([c, ctp + cfn, ctp, cfp, cfn, cp, cr, cf, conf, sum(c in r["unsure"] for r in labelled)])
        tables.append(("Per category (Jev)",
                       ["Category", "Expected", "TP", "FP", "FN", "Precision", "Recall", "F1", "Avg % sure", "Unsure"], rows))

        bands = {"95-100% sure": [0, 0], "65-95% sure": [0, 0], "under 65% sure": [0, 0]}
        for r in labelled:
            for c in CATEGORIES:
                conf = r["categories"][c]["confidence"]
                b = "95-100% sure" if conf >= 0.95 else "65-95% sure" if conf >= 0.65 else "under 65% sure"
                bands[b][0] += 1
                bands[b][1] += r["correct"][c]
        tables.append(("Confidence check (all 7 Jev answers per chat)", ["How sure Jev was", "Answers", "Accuracy"],
                       [[b, n, k / n if n else nan] for b, (n, k) in bands.items()]))

        g_hit = g_n = p_hit = p_n = l_ok = l_n = 0
        for r in labelled:
            gh, ph, lo = span_match(r["text"], r["extracted"], r["labels"]["spans"])
            g_hit, g_n = g_hit + sum(gh), g_n + len(gh)
            p_hit, p_n = p_hit + sum(ph), p_n + len(ph)
            l_ok, l_n = l_ok + sum(lo), l_n + len(lo)
        tables.append(("Extraction (Haiku) vs expected spans",
                       ["Span recall", "Span precision", "Label accuracy", "Expected spans found", "Extracted items correct"],
                       [[g_hit / g_n if g_n else nan, p_hit / p_n if p_n else nan, l_ok / l_n if l_n else nan,
                         f"{g_hit}/{g_n}", f"{p_hit}/{p_n}"]]))

        tag_n, tag_wrong = Counter(), Counter()
        for r in labelled:
            for t in r["tags"]:
                tag_n[t] += 1
                tag_wrong[t] += not r["correct_all"]
        tables.append(("By trap tag", ["Tag", "Chats", "Not fully right"],
                       [[t, n, tag_wrong[t]] for t, n in sorted(tag_n.items(), key=lambda kv: (-tag_wrong[kv[0]], kv[0]))]))

    if ok:
        llm = [r for r in ok if r["llm_called"]]
        med = lambda xs: sorted(xs)[len(xs) // 2] if xs else 0
        jc, lc = sum(r["jev_cost"] for r in ok), sum(r["llm_cost"] for r in llm)
        tables.append(("Cost and speed", ["Model", "Calls", "Input tokens", "Output tokens", "Cost $", "Median ms"], [
            ["Jev", len(ok), sum(r["jev_tokens"] for r in ok), 0, round(jc, 6), round(med([r["jev_ms"] for r in ok]))],
            ["Claude Haiku 4.5", len(llm), sum(r["llm_tokens_in"] for r in llm), sum(r["llm_tokens_out"] for r in llm),
             round(lc, 6), round(med([r["llm_ms"] for r in llm]))],
            ["Total", len(ok) + len(llm), "", "", round(jc + lc, 6), ""]]))
        versions = {r["jev_model_version"] for r in ok}
        if versions != {JEV_PRICED_VERSION}:
            tables.append((f"NOTE: Jev answered as {sorted(versions)}; costs assume {JEV_PRICED_VERSION} pricing. "
                           "Check docs.typesafe.ai/models.md and update JEV_PRICE_INPUT.", None, []))
    if failed:
        tables.append((f"{failed} chats FAILED (see the Results sheet)", None, []))

    # print the same tables (the long trap-tag table is only in the Summary sheet)
    print(f"\n{len(ok)} chats scored" + (f", {failed} FAILED" if failed else "") + (f" in {wall:.1f}s" if wall else ""))
    for title, cols, rows in tables:
        if title == "By trap tag":
            continue
        print(f"\n{title}")
        if cols:
            print("  " + " | ".join(str(c) for c in cols))
        for row in rows:
            print("  " + " | ".join((f"{v:.6f}" if 0 < abs(v) < 0.1 else f"{v:.2f}") if isinstance(v, float) else str(v)
                                    for v in row))
    bad = [r for r in labelled if not r["correct_all"]]
    if bad:
        print(f"\nNot fully right ({len(bad)} chats): " + ", ".join(r["id"] for r in bad))
    return tables


# ---------------------------------------------------------------- later: redaction (not run by default)
def redact_pii(text, extracted):
    """Replace every occurrence of each extracted item with [LABEL]. Longest items first."""
    for it in sorted(extracted, key=lambda i: -len(i["text"])):
        text = text.replace(it["text"], f"[{it['label']}]")
    return text


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", nargs="?", default="ecom_realistic_dev_150.csv")
    ap.add_argument("--limit", type=int, help="only the first N chats")
    ap.add_argument("--ids", help="comma-separated chat ids to run")
    ap.add_argument("--workers", type=int, default=10)
    ap.add_argument("--project", default=None, help="LangSmith project (default: LANGSMITH_PROJECT or pii-detection-with-jev)")
    ap.add_argument("--questions", choices=sorted(QUESTION_SETS), default="v2", help="Jev question wording (default: v2)")
    ap.add_argument("--no-trace", action="store_true", help="don't send traces to LangSmith")
    cfg = ap.parse_args()

    cfg.typesafe_key = env_key("TYPESAFE_API_KEY")
    cfg.openrouter_key = env_key("OPENROUTER_API_KEY")
    if not cfg.typesafe_key:
        sys.exit("TYPESAFE_API_KEY is not set.")
    if not cfg.openrouter_key:
        sys.exit("OPENROUTER_API_KEY is not set.")
    cfg.batch_id = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    cfg.print_lock = threading.Lock()

    # LangSmith: must be configured before the first @traceable call
    cfg.ls_client = None
    if cfg.no_trace:
        os.environ["LANGSMITH_TRACING"] = "false"
    elif env_key("LANGSMITH_API_KEY"):
        os.environ["LANGSMITH_TRACING"] = "true"
        os.environ["LANGSMITH_PROJECT"] = cfg.project or env_key("LANGSMITH_PROJECT") or "pii-detection-with-jev"
        env_key("LANGSMITH_ENDPOINT")  # only needed for EU / self-hosted
        from langsmith import Client
        from langsmith.utils import LangSmithNotFoundError
        cfg.ls_client = Client()
        try:
            cfg.ls_project_id = cfg.ls_client.read_project(project_name=os.environ["LANGSMITH_PROJECT"]).id
        except LangSmithNotFoundError:
            cfg.ls_project_id = cfg.ls_client.create_project(os.environ["LANGSMITH_PROJECT"]).id
    else:
        os.environ["LANGSMITH_TRACING"] = "false"
        print("WARNING: LANGSMITH_API_KEY is not set, so nothing will be traced. Results are still saved locally.")

    ids = set(cfg.ids.split(",")) if cfg.ids else None
    run(cfg.file, cfg, limit=cfg.limit, ids=ids)


if __name__ == "__main__":
    main()
